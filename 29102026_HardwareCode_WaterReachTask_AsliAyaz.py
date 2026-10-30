import time
import random
import datetime
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from tinkerforge.ip_connection import IPConnection
from tinkerforge.bricklet_piezo_speaker_v2 import BrickletPiezoSpeakerV2
from tinkerforge.bricklet_industrial_quad_relay_v2 import BrickletIndustrialQuadRelayV2
from tinkerforge.bricklet_multi_touch_v2 import BrickletMultiTouchV2

HOST = "10.0.0.1"  #IP Adresss
PORT = 4223       #SSID/Hostname: esp32-226Z, and password is in gdrive

UID_PIEZO = "24Gg"
UID_RELAY = "2io5"
UID_TOUCH = "LHs"

TOTAL_TRIALS = 50
TRIAL_DUR    = 15.0
LED_START    = 3.0
LED_DUR      = 6.0
VALVE_TIME   = 6.0
VALVE_DUR    = 0.3

timestamp  = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
DOWNLOADS  = os.path.join(os.path.expanduser("~"), "Downloads")
EXCEL_FILE = os.path.join(DOWNLOADS, f"deney_{timestamp}.xlsx")

# Excel workbook hazırla
wb = Workbook()

# Sheet 1: Trial Data
ws_data = wb.active
ws_data.title = "Trial Data"
headers = ["Trial", "LED", "Licked", "Lick Count", "Trial Duration (s)", "Log"]
ws_data.append(headers)
for col, header in enumerate(headers, 1):
    cell = ws_data.cell(row=1, column=col)
    cell.font      = Font(bold=True, color="FFFFFF")
    cell.fill      = PatternFill("solid", fgColor="2E75B6")
    cell.alignment = Alignment(horizontal="center")

# Sheet 2: Summary
ws_sum = wb.create_sheet("Summary")

# Sheet 3: Log
ws_log = wb.create_sheet("Log")
ws_log.append(["Timestamp", "Event"])
ws_log.cell(row=1, column=1).font = Font(bold=True)
ws_log.cell(row=1, column=2).font = Font(bold=True)

ipcon = IPConnection()
ipcon.set_timeout(2)
piezo = None
relay = None
touch = None

def connect():
    global piezo, relay, touch
    while True:
        try:
            print("  Connecting...")
            try:
                ipcon.disconnect()
            except:
                pass
            time.sleep(1)
            ipcon.connect(HOST, PORT)
            piezo = BrickletPiezoSpeakerV2(UID_PIEZO, ipcon)
            relay = BrickletIndustrialQuadRelayV2(UID_RELAY, ipcon)
            touch = BrickletMultiTouchV2(UID_TOUCH, ipcon)
            touch.set_electrode_config([True, False, False, False, False, False,
                                         False, False, False, False, False, False, False])
            print("  Connected.")
            return
        except Exception as e:
            print(f"  Connection error: {e}")
            print("  Retrying in 5 seconds...")
            time.sleep(5)

def safe_relay(value):
    while True:
        try:
            relay.set_value(value)
            return
        except:
            print("  Relay disconnected, reconnecting...")
            connect()

def all_off():
    safe_relay([False, False, False, False])

def read_touch():
    try:
        state = touch.get_touch_state()
        return state[0]
    except:
        return False

trial_start_ref = [0]
trial_log_lines = []

def log(msg):
    print(msg)
    trial_log_lines.append(msg)

def wait_and_touch(duration, last_touch, first_lick, lick_count):
    end = time.time() + duration
    while time.time() < end:
        elapsed   = time.time() - trial_start_ref[0]
        electrode = read_touch()
        if electrode and not last_touch:
            lick_count += 1
            if not first_lick:
                first_lick = True
                log(f"  [{elapsed:.1f}s] LICK #{lick_count} (first)")
            else:
                log(f"  [{elapsed:.1f}s] LICK #{lick_count}")
        last_touch = electrode
        time.sleep(0.05)
    return last_touch, first_lick, lick_count

def save_excel():
    wb.save(EXCEL_FILE)

def run_trial(trial_no):
    trial_log_lines.clear()
    is_blue  = random.choice([True, False])
    led_name = "BLUE" if is_blue else "GREEN"

    log(f"--- TRIAL {trial_no}/{TOTAL_TRIALS} | LED: {led_name} ---")

    trial_start_ref[0] = time.time()
    last_touch = False
    first_lick = False
    lick_count = 0

    # 0s — Beep
    try:
        piezo.set_beep(1000, 5, 500)
    except:
        pass
    log("  [0.0s] Beep")

    # 0s -> 3s
    last_touch, first_lick, lick_count = wait_and_touch(
        LED_START, last_touch, first_lick, lick_count)

    # 3s — LED on
    if is_blue:
        safe_relay([False, True, False, False])
    else:
        safe_relay([False, False, False, True])
    log(f"  [3.0s] {led_name} LED on")

    # 3s -> 6s
    last_touch, first_lick, lick_count = wait_and_touch(
        VALVE_TIME - LED_START, last_touch, first_lick, lick_count)

    # 6s — Valve open
    if is_blue:
        safe_relay([False, True, True, False])
    else:
        safe_relay([False, False, True, True])
    elapsed = time.time() - trial_start_ref[0]
    log(f"  [{elapsed:.1f}s] Valve open (300ms)")
    time.sleep(VALVE_DUR)

    # Valve close, LED stays
    if is_blue:
        safe_relay([False, True, False, False])
    else:
        safe_relay([False, False, False, True])
    elapsed = time.time() - trial_start_ref[0]
    log(f"  [{elapsed:.1f}s] Valve closed")

    # 6s -> 9s
    last_touch, first_lick, lick_count = wait_and_touch(
        (LED_START + LED_DUR) - VALVE_TIME - VALVE_DUR,
        last_touch, first_lick, lick_count)

    # 9s — LED off
    all_off()
    elapsed = time.time() - trial_start_ref[0]
    log(f"  [{elapsed:.1f}s] LED off")

    # 9s -> 15s
    last_touch, first_lick, lick_count = wait_and_touch(
        TRIAL_DUR - (LED_START + LED_DUR),
        last_touch, first_lick, lick_count)

    all_off()
    duration = time.time() - trial_start_ref[0]
    log(f"  Trial done | Licked: {'YES' if first_lick else 'NO'} | "
        f"Licks: {lick_count} | LED: {led_name}")

    # Trial Data sheet'e ekle
    trial_log_text = "\n".join(trial_log_lines)
    row = [trial_no, led_name, 1 if first_lick else 0, lick_count, round(duration, 1), trial_log_text]
    ws_data.append(row)

    # Renk — LED'e göre
    fill_color = "BDD7EE" if is_blue else "E2EFDA"
    for col in range(1, 6):
        ws_data.cell(row=trial_no+1, column=col).fill = PatternFill("solid", fgColor=fill_color)
        ws_data.cell(row=trial_no+1, column=col).alignment = Alignment(horizontal="center")

    # Log sheet'e ekle
    for line in trial_log_lines:
        ws_log.append([datetime.datetime.now().strftime("%H:%M:%S"), line])

    # Her trial sonrası kaydet
    save_excel()

    return first_lick, led_name, lick_count

def write_summary(completed, licked_total, total_licks,
                  blue_total, blue_licked, blue_licks,
                  green_total, green_licked, green_licks):
    ws_sum.append(["EXPERIMENT SUMMARY"])
    ws_sum["A1"].font = Font(bold=True, size=14)
    ws_sum.append([])
    ws_sum.append(["Parameter", "Value"])
    ws_sum["A3"].font = Font(bold=True)
    ws_sum["B3"].font = Font(bold=True)

    rows = [
        ["Date", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["Completed Trials", f"{completed}/{TOTAL_TRIALS}"],
        ["Total Licks", total_licks],
        ["Total Licked", licked_total],
        [],
        ["BLUE LED", ""],
        ["Blue Trials", blue_total],
        ["Blue Licked", f"{blue_licked}/{blue_total} ({blue_licked/blue_total*100:.1f}%)" if blue_total > 0 else "-"],
        ["Blue Licks", blue_licks],
        [],
        ["GREEN LED", ""],
        ["Green Trials", green_total],
        ["Green Licked", f"{green_licked}/{green_total} ({green_licked/green_total*100:.1f}%)" if green_total > 0 else "-"],
        ["Green Licks", green_licks],
        [],
        ["Overall Success", f"{licked_total/completed*100:.1f}%" if completed > 0 else "-"],
    ]
    for row in rows:
        ws_sum.append(row)

    ws_sum.column_dimensions['A'].width = 25
    ws_sum.column_dimensions['B'].width = 30
    ws_data.column_dimensions['A'].width = 8
    ws_data.column_dimensions['B'].width = 8
    ws_data.column_dimensions['C'].width = 8
    ws_data.column_dimensions['D'].width = 12
    ws_data.column_dimensions['E'].width = 20
    ws_data.column_dimensions['F'].width = 60
    ws_log.column_dimensions['A'].width  = 12
    ws_log.column_dimensions['B'].width  = 60
    save_excel()

# Main
print("=" * 40)
print("Water Reaching Task — denemev2")
print(f"Total trials : {TOTAL_TRIALS}")
print(f"Trial dur    : {TRIAL_DUR}s")
print(f"LED start    : {LED_START}s")
print(f"LED dur      : {LED_DUR}s")
print(f"Valve time   : {VALVE_TIME}s")
print(f"Valve dur    : {VALVE_DUR*1000:.0f}ms")
print("=" * 40)

connect()
save_excel()

print(f"Excel: {EXCEL_FILE}")
print("Press Enter to start, Ctrl+C to stop")
input()

licked_total = 0
blue_licked  = 0
green_licked = 0
blue_total   = 0
green_total  = 0
total_licks  = 0
blue_licks   = 0
green_licks  = 0
completed    = 0

try:
    for i in range(1, TOTAL_TRIALS + 1):
        licked, led, licks = run_trial(i)

        if led == "BLUE":
            blue_total += 1
            blue_licks += licks
            if licked:
                blue_licked += 1
        else:
            green_total += 1
            green_licks += licks
            if licked:
                green_licked += 1

        if licked:
            licked_total += 1
        total_licks += licks
        completed = i

        if i < TOTAL_TRIALS:
            time.sleep(0.5)

    print("\n--- ALL TRIALS COMPLETED ---")

except KeyboardInterrupt:
    all_off()
    print("\n--- EXPERIMENT STOPPED ---")

finally:
    write_summary(completed, licked_total, total_licks,
                  blue_total, blue_licked, blue_licks,
                  green_total, green_licked, green_licks)
    print(f"\nExcel saved: {EXCEL_FILE}")
    try:
        ipcon.disconnect()
    except:
        pass
    print("System closed.")